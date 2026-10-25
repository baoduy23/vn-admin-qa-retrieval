"""Đọc qrels từ CSV hoặc phiếu gán nhãn .xlsx (dùng chung cho kappa.py).

CSV  : cột question_id, chunk_id, relevance (như data/labels/qrels.csv).
.xlsx: sheet "qrels" cùng 3 cột đó; cần cài openpyxl.
Trả về list dict, giá trị đã ép về chuỗi và bỏ khoảng trắng ("2.0" từ Excel thành "2").
"""
from __future__ import annotations

import csv
from pathlib import Path

COLS = ("question_id", "chunk_id", "relevance")


def _clean(row: dict) -> dict:
    out = {k: str(row.get(k) if row.get(k) is not None else "").strip() for k in COLS}
    if out["relevance"].endswith(".0"):  # Excel lưu số thành float
        out["relevance"] = out["relevance"][:-2]
    return out


def read_qrels(path: Path) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        if "qrels" not in wb.sheetnames:
            raise SystemExit(f"{path}: thiếu sheet 'qrels'")
        rows = list(wb["qrels"].iter_rows(values_only=True))
        head = [str(h).strip() if h is not None else "" for h in rows[0]]
        data = [dict(zip(head, r)) for r in rows[1:]]
    else:
        with open(path, encoding="utf-8-sig", newline="") as f:
            data = list(csv.DictReader(f))
    if data and not set(COLS) <= set(data[0]):
        raise SystemExit(f"{path}: cần các cột {COLS}, đang có {sorted(data[0])}")
    return [_clean(r) for r in data if r.get("question_id")]
