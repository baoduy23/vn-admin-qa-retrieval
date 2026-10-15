"""Cohen's kappa giữa hai người gán nhãn độc lập (phần 20% nhóm).

    python -m src.labels.kappa --a phieu_A.xlsx --b phieu_B.xlsx \
        --questions data/labels/questions.csv --chunks data/processed/chunks.jsonl \
        --out results/data_quality/label_agreement

Mỗi file là CSV (question_id, chunk_id, relevance) hoặc phiếu .xlsx (sheet qrels, Kết luận).
Câu được tính: giao của hai bên (sheet Kết luận có verdict ở cả hai phiếu; với CSV thì
dùng --qids hoặc các câu có nhãn ở ít nhất một bên).

Đơn vị: cặp (câu có đáp án, chunk) trong tập ứng viên = chunk một trong hai người đã gán
+ mọi chunk cùng Điều với các chunk đó (tránh kappa cao giả vì quá nhiều cặp không liên quan).
Báo cáo: kappa nhị phân (liên quan: mức >=1), kappa có trọng số 3 mức (0/1/2),
kappa verdict ở mức câu hỏi (có/không có đáp án), tỷ lệ đồng ý về type.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from pathlib import Path

from src.labels.validate_labels import read_qrels

REFUSE = {"out_of_scope", "unsupported"}


def cohen(x: list, y: list) -> float:
    n = len(x)
    if n == 0:
        return float("nan")
    cats = sorted(set(x) | set(y))
    po = sum(a == b for a, b in zip(x, y)) / n
    pe = sum((x.count(c) / n) * (y.count(c) / n) for c in cats)
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def weighted_kappa(x: list[int], y: list[int], k: int = 3) -> float:
    n = len(x)
    if n == 0:
        return float("nan")
    w = lambda i, j: 1 - abs(i - j) / (k - 1)
    po = sum(w(a, b) for a, b in zip(x, y)) / n
    px = [x.count(i) / n for i in range(k)]
    py = [y.count(i) / n for i in range(k)]
    pe = sum(w(i, j) * px[i] * py[j] for i in range(k) for j in range(k))
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def read_verdicts(path: Path) -> dict[str, dict]:
    if path.suffix.lower() != ".xlsx":
        return {}
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    if "Kết luận" not in wb.sheetnames:
        return {}
    rows = list(wb["Kết luận"].iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    out = {}
    for r in rows[1:]:
        d = dict(zip(head, r))
        if d.get("question_id") and d.get("verdict"):
            out[str(d["question_id"]).strip()] = d
    return out


def fmt(x: float) -> str:
    return "n/a" if isinstance(x, float) and math.isnan(x) else f"{x:.3f}"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--a", type=Path, required=True)
    ap.add_argument("--b", type=Path, required=True)
    ap.add_argument("--questions", type=Path, default=Path("data/labels/questions.csv"))
    ap.add_argument("--chunks", type=Path, default=Path("data/processed/chunks.jsonl"))
    ap.add_argument("--qids", type=Path, help="file mỗi dòng một question_id (khi dùng CSV)")
    ap.add_argument("--out", type=Path)
    a = ap.parse_args()

    with open(a.questions, encoding="utf-8-sig", newline="") as f:
        qs = {r["question_id"].strip(): r for r in csv.DictReader(f)}
    meta = {}
    for l in open(a.chunks, encoding="utf-8"):
        c = json.loads(l)
        meta[c["chunk_id"]] = (c["doc_id"], c["dieu"])
    by_dieu = defaultdict(set)
    for cid, k in meta.items():
        by_dieu[k].add(cid)

    def labels(path):
        d = defaultdict(dict)
        for r in read_qrels(path):
            if r["relevance"] in ("1", "2"):
                d[r["question_id"]][r["chunk_id"]] = int(r["relevance"])
        return d

    la, lb = labels(a.a), labels(a.b)
    va, vb = read_verdicts(a.a), read_verdicts(a.b)
    if a.qids:
        qids = [l.strip() for l in open(a.qids, encoding="utf-8") if l.strip()]
    elif va and vb:
        qids = sorted(set(va) & set(vb))
    else:
        qids = sorted((set(la) | set(lb)) & set(qs))
    qids = [q for q in qids if q in qs]

    bx, by, wx, wy, diffs = [], [], [], [], []
    for q in qids:
        if qs[q]["type"] in REFUSE:
            continue
        sel = set(la.get(q, {})) | set(lb.get(q, {}))
        pool = set(sel)
        for c in sel:
            pool |= by_dieu.get(meta.get(c), set())
        for c in sorted(pool):
            ra, rb = la.get(q, {}).get(c, 0), lb.get(q, {}).get(c, 0)
            bx.append(int(ra >= 1)); by.append(int(rb >= 1)); wx.append(ra); wy.append(rb)
            if ra != rb:
                diffs.append((q, c, ra, rb))

    def verdict(q, v, l):
        if q in v:
            return v[q]["verdict"]
        return "answerable" if l.get(q) else "no_answer"
    ya = [verdict(q, va, la) for q in qids]
    yb = [verdict(q, vb, lb) for q in qids]
    vdiff = [q for q, x, y in zip(qids, ya, yb) if x != y]
    type_rate = {}
    for name, v in (("A", va), ("B", vb)):
        got = [d.get("type_agree") for d in v.values() if q_in(d, qids)] if v else []
        type_rate[name] = (sum(g == "yes" for g in got) / len(got)) if got else float("nan")

    kb, kw, kv = cohen(bx, by), weighted_kappa(wx, wy), cohen(ya, yb)
    lines = [f"# Độ đồng thuận gán nhãn", "",
             f"- Số câu: {len(qids)} ({sum(qs[q]['type'] not in REFUSE for q in qids)} câu có đáp án)",
             f"- Số cặp (câu, chunk) trong tập ứng viên: {len(bx)}",
             f"- **Cohen's kappa (liên quan / không liên quan): {fmt(kb)}**",
             f"- Kappa có trọng số 3 mức (0/1/2): {fmt(kw)}",
             f"- Kappa verdict mức câu (có/không có đáp án): {fmt(kv)}",
             f"- Đồng ý về type: người A {fmt(type_rate['A'])}, người B {fmt(type_rate['B'])}",
             f"- Số cặp lệch nhãn: {len(diffs)}; số câu lệch verdict: {len(vdiff)}", ""]
    if not math.isnan(kb) and kb < 0.6:
        lines += ["**Kappa < 0,6: rà lại guideline (ranh giới mức 1/2, ranh giới loại câu) trước khi gán lại phần lệch.**", ""]
    lines += ["## Cặp lệch (họp chốt)", "", "| question_id | chunk_id | A | B |", "|---|---|---|---|"]
    lines += [f"| {q} | {c} | {x} | {y} |" for q, c, x, y in diffs]
    if vdiff:
        lines += ["", "## Câu lệch verdict", ""] + [f"- {q}" for q in vdiff]
    text = "\n".join(lines) + "\n"
    if a.out:
        a.out.parent.mkdir(parents=True, exist_ok=True)
        a.out.with_suffix(".md").write_text(text, encoding="utf-8")
        with open(a.out.with_name(a.out.name + "_diffs.csv"), "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(["question_id", "chunk_id", "rel_A", "rel_B"])
            w.writerows(diffs)
    print(text)


def q_in(d: dict, qids: list[str]) -> bool:
    return str(d.get("question_id", "")).strip() in qids


if __name__ == "__main__":
    main()
